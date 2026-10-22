"""Importação de listas de contatos (Excel, CSV, Word ou texto colado) para prospecção.

A lista vira uma "busca" de origem importação: cada contato é guardado como empresa e passa pelo
mesmo processamento das buscas do Apify (blacklist, já era lead, conferência de WhatsApp). Assim os
leads da lista aparecem nas campanhas como os de qualquer busca.

As funções alteram a sessão mas não fazem commit.
"""

import csv
import io
import logging
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field

from sqlalchemy.orm import Session

from app.models import Busca, BuscaStatus, Campanha, CampanhaStatus, Empresa, User
from app.services.telefone import eh_celular, normalizar_telefone

logger = logging.getLogger(__name__)

TAMANHO_MAXIMO = 10 * 1024 * 1024  # 10 MB
MAXIMO_CONTATOS = 20000
EXTENSOES = (".xlsx", ".xls", ".csv", ".txt", ".docx")

# Sequência que pode ser um telefone: começa e termina com dígito, com espaços, pontos, traços e parênteses no meio.
_TELEFONE = re.compile(r"\+?\(?\d[\d\s().\-]{7,18}\d")
_SEPARADORES = " \t-–—:;,|/•·"

# Palavras dos cabeçalhos, sem acento e em minúsculas.
COLUNAS = {
    "telefone": ("telefone", "fone", "celular", "whatsapp", "whats", "zap", "numero", "contato", "tel", "phone", "mobile"),
    "nome": ("nome", "empresa", "razao", "fantasia", "cliente", "estabelecimento", "loja", "title", "name", "company"),
    "cidade": ("cidade", "municipio", "city"),
    "categoria": ("segmento", "categoria", "ramo", "atividade", "nicho", "category"),
}


# Cabeçalhos exatos de exportações conhecidas (sem acento, só letras e números): extratores do Google Maps
# (OrganizationName, OrganizationPhoneNr…) e a planilha do Apify (title, phone, city…).
CABECALHOS_CONHECIDOS = {
    "organizationphonenr": "telefone", "phoneunformatted": "telefone",
    "organizationphonefromthewebsite": "telefone_alt",
    "organizationname": "nome", "title": "nome",
    "organizationcategory": "categoria", "categoryname": "categoria",
    "organizationaddress": "endereco", "address": "endereco", "endereco": "endereco",
    "state": "estado", "estado": "estado", "uf": "estado",
    "organizationwebsite": "site", "website": "site", "site": "site",
    "organizationemail": "email", "email": "email", "emails": "email",
    "ratestars": "nota", "totalscore": "nota", "nota": "nota",
    "reviewstotalcount": "avaliacoes", "reviewscount": "avaliacoes", "avaliacoes": "avaliacoes",
    "gmapslink": "url_maps", "url": "url_maps",
    "organizationgmapsid": "place_id", "placeid": "place_id",
    "searchrequest": "busca",
    "organizationinstagram": "instagram", "instagram": "instagram",
    "organizationfacebook": "facebook", "facebook": "facebook",
    "neighborhood": "bairro", "bairro": "bairro",
}
# "Endereço: R. X, 10 - Centro, Angra dos Reis - RJ, 23900-000" -> cidade e UF
_CIDADE_UF = re.compile(r"(?:^|,)\s*([^,]+?)\s+-\s+([A-Z]{2})(?:,\s*\d{5}-?\d{3})?\s*$")


class ImportacaoError(Exception):
    pass


@dataclass
class Contato:
    linha: int
    telefone_original: str
    nome: str = ""
    cidade: str = ""
    categoria: str = ""
    estado: str = ""
    bairro: str = ""
    endereco: str = ""
    site: str = ""
    email: str = ""
    nota: float | None = None
    avaliacoes: int | None = None
    url_maps: str = ""
    place_id: str = ""
    extras: dict = field(default_factory=dict)  # instagram, facebook, busca de origem


def _sem_acento(texto: str) -> str:
    return unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode().lower().strip()


def _celula(valor) -> str:
    """O Excel guarda telefone como número (11999998888.0); devolve sempre texto."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


# --- Leitura dos arquivos: tudo vira uma lista de linhas (listas de células) -------------


def _decodificar(conteudo: bytes) -> str:
    for codificacao in ("utf-8-sig", "cp1252"):
        try:
            return conteudo.decode(codificacao)
        except UnicodeDecodeError:
            continue
    return conteudo.decode("latin-1")


def _ler_csv(conteudo: bytes) -> list[list[str]]:
    texto = _decodificar(conteudo)
    try:
        dialeto = csv.Sniffer().sniff(texto[:5000], delimiters=";,\t|")
    except csv.Error:
        return [[linha] for linha in texto.splitlines()]  # sem separador: uma coluna só
    return [[_celula(c) for c in linha] for linha in csv.reader(io.StringIO(texto), dialeto)]


def _ler_xlsx(conteudo: bytes) -> list[list[list[str]]]:
    import openpyxl

    try:
        livro = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise ImportacaoError("Não consegui abrir a planilha .xlsx. Confira se o arquivo não está corrompido.") from exc
    abas = [[[_celula(c) for c in linha] for linha in aba.iter_rows(values_only=True)] for aba in livro.worksheets]
    livro.close()
    return abas


def _ler_xls(conteudo: bytes) -> list[list[list[str]]]:
    import xlrd

    try:
        livro = xlrd.open_workbook(file_contents=conteudo)
    except xlrd.XLRDError as exc:
        raise ImportacaoError("Não consegui abrir a planilha .xls. Tente salvar como .xlsx ou .csv.") from exc
    return [[[_celula(c) for c in aba.row_values(i)] for i in range(aba.nrows)] for aba in livro.sheets()]


def _ler_docx(conteudo: bytes) -> list[list[str]]:
    import docx

    try:
        documento = docx.Document(io.BytesIO(conteudo))
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ImportacaoError("Não consegui abrir o documento .docx.") from exc
    linhas = [[p.text] for p in documento.paragraphs if p.text.strip()]
    for tabela in documento.tables:
        linhas.extend([_celula(c.text) for c in linha.cells] for linha in tabela.rows)
    return linhas


def ler_abas(nome: str, conteudo: bytes) -> list[list[list[str]]]:
    """Cada aba da planilha separada: cada uma tem o próprio cabeçalho (ex.: "Organizations" e "Reviews")."""
    nome = (nome or "").lower()
    if len(conteudo) > TAMANHO_MAXIMO:
        raise ImportacaoError("O arquivo passou de 10 MB. Divida a lista em partes menores.")
    if nome.endswith(".doc"):
        raise ImportacaoError("Arquivo .doc antigo não é aceito: abra no Word e salve como .docx, ou cole os números no campo de texto.")
    if nome.endswith(".xlsx"):
        return _ler_xlsx(conteudo)
    if nome.endswith(".xls"):
        return _ler_xls(conteudo)
    if nome.endswith(".docx"):
        return [_ler_docx(conteudo)]
    if nome.endswith((".csv", ".txt")):
        return [_ler_csv(conteudo)]
    raise ImportacaoError("Formato não aceito. Use Excel (.xlsx ou .xls), CSV, TXT ou Word (.docx).")


def ler_arquivo(nome: str, conteudo: bytes) -> list[list[str]]:
    """As linhas da primeira aba (para arquivos de uma aba só)."""
    abas = ler_abas(nome, conteudo)
    return abas[0] if abas else []


def contatos_do_arquivo(nome: str, conteudo: bytes) -> list["Contato"]:
    return [contato for aba in ler_abas(nome, conteudo) for contato in extrair_contatos(aba)]


def ler_texto(texto: str) -> list[list[str]]:
    """Texto colado: uma linha por contato, com o telefone e, se tiver, o nome."""
    linhas = (texto or "").splitlines()
    if any("\t" in linha for linha in linhas):  # colado do Excel: as colunas vêm separadas por tab
        return [[c.strip() for c in linha.split("\t")] for linha in linhas]
    return [[linha] for linha in linhas]


# --- Extração dos contatos -----------------------------------------------------


def _mapear_cabecalho(linha: list[str]) -> dict[str, int] | None:
    """Se a linha é um cabeçalho com uma coluna de telefone, devolve campo -> índice da coluna."""
    if sum(1 for celula in linha if celula.strip()) < 2:
        return None  # numa coluna só, "linha sem número" não é cabeçalho
    mapa: dict[str, int] = {}
    for indice, celula in enumerate(linha):
        texto = _sem_acento(celula)
        if not texto:
            continue
        conhecido = CABECALHOS_CONHECIDOS.get(re.sub(r"[^a-z0-9]", "", texto))
        if conhecido:
            mapa.setdefault(conhecido, indice)
            continue
        # Vale a primeira palavra conhecida: "Nome do contato" é nome, "Telefone da empresa" é telefone.
        for palavra in re.split(r"[^a-z0-9]+", texto):
            campo = next((c for c, chaves in COLUNAS.items() if palavra in chaves), None)
            if campo:
                mapa.setdefault(campo, indice)
                break
    return mapa if "telefone" in mapa else None


def telefones_no_texto(texto: str) -> list[str]:
    return [m.group(0).strip() for m in _TELEFONE.finditer(texto or "") if len(re.sub(r"\D", "", m.group(0))) >= 10]


def _nome_da_linha(celulas: list[str], telefones: list[str]) -> str:
    """O primeiro texto com letras que não é telefone (numa linha de uma célula, o que sobra sem o telefone)."""
    for celula in celulas:
        resto = celula
        for telefone in telefones:
            resto = resto.replace(telefone, " ")
        resto = re.sub(r"\s{2,}", " ", resto).strip(_SEPARADORES + " ")
        if re.search(r"[A-Za-zÀ-ÿ]", resto):
            return resto[:300]
    return ""


def _numero(texto: str, inteiro: bool = False):
    texto = (texto or "").strip().replace(",", ".")
    try:
        return int(float(texto)) if inteiro else float(texto)
    except ValueError:
        return None


def _primeiro_email(texto: str) -> str:
    """A coluna de e-mail dos extratores traz vários, inclusive de serviços do site (sentry, wix)."""
    for parte in re.split(r"[,;\s]+", texto or ""):
        if "@" in parte and not any(lixo in parte for lixo in ("sentry", "wixpress", "example.")):
            return parte[:255]
    return ""


def cidade_do_endereco(endereco: str) -> tuple[str, str, str]:
    """(endereço limpo, cidade, UF) a partir de "Endereço: ..., Cidade - UF, CEP"."""
    limpo = re.sub(r"^\s*endere\S*\s*:\s*", "", endereco or "", flags=re.IGNORECASE).strip()
    achado = _CIDADE_UF.search(limpo)
    return limpo[:500], (achado.group(1).strip()[:200] if achado else ""), (achado.group(2) if achado else "")


def _contato_da_linha(numero_linha: int, telefone: str, pegar) -> Contato:
    endereco, cidade_end, uf_end = cidade_do_endereco(pegar("endereco"))
    busca = pegar("busca").replace("+", " ").strip()
    extras = {k: v for k in ("instagram", "facebook") if (v := pegar(k).split(",")[0].strip())}
    if busca:
        extras["busca"] = busca
    return Contato(
        linha=numero_linha,
        telefone_original=telefone[:50],
        nome=pegar("nome")[:300],
        cidade=(pegar("cidade") or cidade_end)[:200],
        categoria=pegar("categoria")[:200],
        estado=(pegar("estado") or uf_end)[:100],
        bairro=pegar("bairro")[:200],
        endereco=endereco,
        site=pegar("site")[:500],
        email=_primeiro_email(pegar("email")),
        nota=_numero(pegar("nota")),
        avaliacoes=_numero(pegar("avaliacoes"), inteiro=True),
        url_maps=pegar("url_maps")[:1000],
        place_id=pegar("place_id")[:200],
        extras=extras,
    )


def extrair_contatos(linhas: list[list[str]]) -> list[Contato]:
    """Acha os telefones nas linhas. Com cabeçalho, usa as colunas; sem, procura em cada célula."""
    linhas = [[_celula(c) for c in linha] for linha in linhas]
    mapa, inicio = None, 0
    for i, linha in enumerate(linhas[:10]):
        if any(telefones_no_texto(celula) for celula in linha):
            break  # o cabeçalho vem antes do primeiro telefone
        mapa = _mapear_cabecalho(linha)
        if mapa:
            inicio = i + 1
            break

    contatos: list[Contato] = []
    for numero_linha, linha in enumerate(linhas[inicio:], start=inicio + 1):
        if not any(linha):
            continue
        if mapa:
            pegar = lambda campo: linha[mapa[campo]] if campo in mapa and mapa[campo] < len(linha) else ""  # noqa: E731
            celula_telefone = pegar("telefone") or pegar("telefone_alt")  # sem o do Maps, o achado no site
            # Sem formato de telefone, só conta como "telefone inválido" se tiver algum número.
            telefones = telefones_no_texto(celula_telefone) or ([celula_telefone] if any(c.isdigit() for c in celula_telefone) else [])
            contatos.extend(_contato_da_linha(numero_linha, telefone, pegar) for telefone in telefones)
        else:
            telefones = [t for celula in linha for t in telefones_no_texto(celula)]
            nome = _nome_da_linha(linha, telefones)
            contatos.extend(Contato(numero_linha, telefone[:50], nome) for telefone in telefones)
    return contatos


# --- Gravação -----------------------------------------------------------------------


def importar(
    db: Session,
    contatos: list[Contato],
    *,
    nome_lista: str,
    cidade: str = "",
    segmento: str = "",
    campanha: Campanha | None = None,
    user: User | None = None,
) -> Busca:
    """Cria a lista (busca de origem importação) com os contatos guardados, pronta para processar."""
    nome_lista = nome_lista.strip()[:200]
    if not nome_lista:
        raise ImportacaoError("Dê um nome para a lista.")
    if not contatos:
        raise ImportacaoError("Não achei nenhum telefone. Confira se a lista tem uma coluna de telefone ou um número por linha.")
    if len(contatos) > MAXIMO_CONTATOS:
        raise ImportacaoError(f"A lista tem {len(contatos)} telefones; o máximo por importação é {MAXIMO_CONTATOS}. Divida em partes.")
    if campanha is not None and campanha.status == CampanhaStatus.ENCERRADA:
        raise ImportacaoError("A campanha escolhida está encerrada.")

    busca = Busca(
        termo=nome_lista,
        cidade=cidade.strip()[:200],
        limite=len(contatos),
        origem="importacao",
        # "executando" é o estado de quem espera o processamento; sem run do Apify, o polling não a toca.
        status=BuscaStatus.EXECUTANDO,
        campanha_id=campanha.id if campanha else None,
        criado_por_id=user.id if user else None,
    )
    db.add(busca)
    db.flush()
    vistos: set[str] = set()
    for contato in contatos:
        if contato.place_id:
            if contato.place_id in vistos:
                continue  # a mesma empresa repetida na planilha
            vistos.add(contato.place_id)
        telefone = normalizar_telefone(contato.telefone_original)
        db.add(
            Empresa(
                busca_id=busca.id,
                nome=contato.nome,
                telefone_original=contato.telefone_original,
                telefone=telefone,
                celular=eh_celular(telefone),
                # Sem categoria nem segmento, a busca que gerou a planilha ("angra dos reis-rj pousada") ajuda a achar o roteiro.
                categoria=(contato.categoria or segmento.strip() or contato.extras.get("busca", ""))[:200] or None,
                cidade=(contato.cidade or cidade.strip())[:200] or None,
                estado=contato.estado or None,
                bairro=contato.bairro or None,
                endereco=contato.endereco or None,
                site=contato.site or None,
                nota=contato.nota,
                avaliacoes=contato.avaliacoes,
                url_maps=contato.url_maps or None,
                place_id=contato.place_id or None,
                dados={"origem": "importacao", "linha": contato.linha, **({"email": contato.email} if contato.email else {}), **contato.extras},
            )
        )
    db.flush()
    logger.info("lista importada", extra={"busca_id": busca.id, "contatos": len(contatos), "campanha_id": busca.campanha_id})
    return busca
