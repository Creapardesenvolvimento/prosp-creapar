import io

import pytest
from sqlalchemy import select

from app.models import Blacklist, Busca, BuscaStatus, Empresa, Lead, LeadStatus
from app.services import buscas, importacao
from app.services.importacao import Contato, ImportacaoError
from tests.test_disparo import criar_campanha, criar_numero


def test_extrai_pelas_colunas_do_cabecalho():
    linhas = [
        ["Lista da associação", "", ""],
        ["Razão social", "Cidade", "WhatsApp", "Segmento", "Nome do contato"],
        ["Padaria Sol", "Campinas", 19999998888.0, "padaria", "João"],
        ["Oficina do Zé", "", "(19) 3333-4444 / 19 97777-6666", "", ""],
        ["", "", "", "", ""],
        ["Sem telefone", "Jundiaí", "", "", ""],
    ]
    contatos = importacao.extrair_contatos(linhas)
    assert [(c.nome, c.telefone_original, c.cidade, c.categoria) for c in contatos] == [
        ("Padaria Sol", "19999998888", "Campinas", "padaria"),
        ("Oficina do Zé", "(19) 3333-4444", "", ""),
        ("Oficina do Zé", "19 97777-6666", "", ""),
    ]
    assert contatos[0].linha == 3


def test_extrai_sem_cabecalho_de_texto_colado():
    texto = "Padaria Sol - (19) 99999-8888\n\n11 98888-7777\nOficina do Zé;19 97777-6666\nlinha sem número\n+55 (21) 96666-5555 Mercado Bom"
    contatos = importacao.extrair_contatos(importacao.ler_texto(texto))
    assert [(c.nome, c.telefone_original) for c in contatos] == [
        ("Padaria Sol", "(19) 99999-8888"),
        ("", "11 98888-7777"),
        ("Oficina do Zé", "19 97777-6666"),
        ("Mercado Bom", "+55 (21) 96666-5555"),
    ]


def test_texto_colado_do_excel_usa_as_colunas():
    contatos = importacao.extrair_contatos(importacao.ler_texto("Empresa\tTelefone\nPadaria Sol\t19 99999-8888"))
    assert [(c.nome, c.telefone_original) for c in contatos] == [("Padaria Sol", "19 99999-8888")]


def test_le_xlsx_csv_e_docx():
    import docx
    import openpyxl

    livro = openpyxl.Workbook()
    livro.active.append(["Nome", "Celular"])
    livro.active.append(["Padaria Sol", 19999998888])
    arquivo = io.BytesIO()
    livro.save(arquivo)
    contatos = importacao.extrair_contatos(importacao.ler_arquivo("lista.XLSX", arquivo.getvalue()))
    assert [(c.nome, c.telefone_original) for c in contatos] == [("Padaria Sol", "19999998888")]

    csv_ = "Empresa;Telefone;Cidade\nMercado Bom;(11) 98888-7777;São Paulo\n".encode("cp1252")
    contatos = importacao.extrair_contatos(importacao.ler_arquivo("lista.csv", csv_))
    assert [(c.nome, c.cidade) for c in contatos] == [("Mercado Bom", "São Paulo")]

    documento = docx.Document()
    documento.add_paragraph("Oficina do Zé — 19 97777-6666")
    tabela = documento.add_table(rows=1, cols=2)
    tabela.rows[0].cells[0].text, tabela.rows[0].cells[1].text = "Bar do Léo", "21 96666-5555"
    arquivo = io.BytesIO()
    documento.save(arquivo)
    contatos = importacao.extrair_contatos(importacao.ler_arquivo("lista.docx", arquivo.getvalue()))
    assert [(c.nome, c.telefone_original) for c in contatos] == [("Oficina do Zé", "19 97777-6666"), ("Bar do Léo", "21 96666-5555")]


@pytest.mark.parametrize(
    "nome, conteudo, erro",
    [
        ("lista.doc", b"x", "salve como .docx"),
        ("lista.pdf", b"x", "Formato não aceito"),
        ("lista.xlsx", b"nao e zip", "Não consegui abrir"),
        ("lista.xls", b"nao e xls", "Não consegui abrir"),
        ("lista.csv", None, "10 MB"),
    ],
    ids=["doc", "pdf", "xlsx-quebrado", "xls-quebrado", "grande"],
)
def test_arquivos_recusados(nome, conteudo, erro):
    with pytest.raises(ImportacaoError, match=erro):
        importacao.ler_arquivo(nome, conteudo if conteudo is not None else b"x" * (importacao.TAMANHO_MAXIMO + 1))


def test_importar_e_processar(db, server, evo):
    criar_numero(db, server, "Chip", 10)
    db.add(Blacklist(telefone="+5519955550002", motivo="opt-out"))
    db.add(Lead(nome="Antigo", telefone="+5519955550003", status=LeadStatus.ENVIADO))
    db.commit()
    evo.sem_whatsapp = {"5519955550004"}
    contatos = [
        Contato(1, "19 95555-0001", "Padaria Sol", "", "padaria"),
        Contato(2, "19 95555-0001", "Padaria Sol repetida"),
        Contato(3, "19 95555-0002", "Na blacklist"),
        Contato(4, "19 95555-0003", "Já é lead"),
        Contato(5, "19 95555-0004", "Sem WhatsApp"),
        Contato(6, "19 3333-4444", "Fixo com WhatsApp Business"),
        Contato(7, "123", "Número quebrado"),
        Contato(8, "11 95555-0008", ""),
    ]
    busca = importacao.importar(db, contatos, nome_lista="Associação", cidade="Campinas", segmento="comércio")
    db.commit()
    assert (busca.origem, busca.status, busca.rotulo) == ("importacao", BuscaStatus.EXECUTANDO, "Lista: Associação")
    assert buscas.buscas_em_andamento(db) == []  # sem run do Apify: o polling não toca nela

    assert buscas.reivindicar_processamento(db, busca.id)
    buscas.processar(db, busca, pausa_entre_lotes=0)
    db.commit()
    assert (busca.status, busca.total_encontrados, busca.total_novos) == (BuscaStatus.CONCLUIDA, 8, 3)
    assert busca.descartes == {
        "repetido_na_busca": 1, "blacklist": 1, "ja_era_lead": 1, "sem_whatsapp": 1, "telefone_invalido": 1,
    }
    leads = {l.telefone: l for l in db.scalars(select(Lead).where(Lead.busca_id == busca.id))}
    assert set(leads) == {"+5519955550001", "+551933334444", "+5511955550008"}
    assert (leads["+5519955550001"].categoria, leads["+5519955550001"].cidade) == ("padaria", "Campinas")
    assert leads["+5511955550008"].nome == "" and leads["+5511955550008"].categoria == "comércio"


def test_lista_vai_direto_para_a_campanha(db, server, evo):
    n = criar_numero(db, server, "Chip", 10)
    c = criar_campanha(db, [n])
    busca = importacao.importar(db, [Contato(1, "19 95555-0001", "Padaria Sol")], nome_lista="Lista", campanha=c)
    db.commit()
    buscas.reivindicar_processamento(db, busca.id)
    buscas.processar(db, busca, pausa_entre_lotes=0)
    db.commit()
    lead = db.scalar(select(Lead).where(Lead.busca_id == busca.id))
    assert (lead.campanha_id, lead.status) == (c.id, LeadStatus.NA_FILA)


def test_importar_valida(db):
    with pytest.raises(ImportacaoError, match="nenhum telefone"):
        importacao.importar(db, [], nome_lista="Lista")
    with pytest.raises(ImportacaoError, match="nome"):
        importacao.importar(db, [Contato(1, "19 95555-0001")], nome_lista=" ")


def test_validacao_reveza_entre_os_numeros(db, server, evo):
    a = criar_numero(db, server, "Chip A", 10)
    b = criar_numero(db, server, "Chip B", 10)
    buscas.validar_whatsapp(db, [f"+55199{i:08d}" for i in range(120)], pausa_entre_lotes=0)
    assert [c[1] for c in evo.called("check_whatsapp")] == [a.instance_name, b.instance_name, a.instance_name]


def test_tela_de_importacao(logged_client, db, server, evo):
    n = criar_numero(db, server, "Chip", 10)
    c = criar_campanha(db, [n])
    assert "Importar lista de contatos" in logged_client.get("/importar").text
    assert "Importar lista" in logged_client.get("/buscas").text

    resp = logged_client.post("/importar", data={"texto": "", "nome": "x"})
    assert resp.status_code == 400 and "Escolha um arquivo" in resp.text

    arquivo = "Empresa;Telefone\nPadaria Sol;19 95555-0001\nMercado;11 95555-0002\n".encode()
    resp = logged_client.post(
        "/importar",
        data={"texto": "Bar do Léo 21 95555-0003", "campanha_id": str(c.id), "cidade": "Campinas"},
        files={"arquivo": ("oficinas.csv", arquivo, "text/csv")},
    )
    assert resp.status_code == 303
    pagina = logged_client.get(resp.headers["location"]).text
    assert "Lista: oficinas" in pagina and "3 telefones lidos" in pagina and "viraram lead" in pagina
    from app.db import SessionLocal

    with SessionLocal() as s:
        busca = s.scalar(select(Busca).where(Busca.origem == "importacao"))
        assert busca.status == BuscaStatus.CONCLUIDA and busca.total_novos == 3
        assert s.scalar(select(Empresa.cidade).where(Empresa.nome == "Bar do Léo")) == "Campinas"
        assert {l.campanha_id for l in s.scalars(select(Lead))} == {c.id}
